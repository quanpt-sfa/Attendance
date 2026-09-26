import sys
try:
    import openpyxl
    wb = openpyxl.load_workbook('dslop.xlsx')
    ws = wb.active
    print("Headers:", [cell.value for cell in ws[1]])
    print("First Row Data:", [cell.value for cell in ws[2]])
except ImportError:
    print("openpyxl not installed")
except Exception as e:
    print(f"Error: {e}")
